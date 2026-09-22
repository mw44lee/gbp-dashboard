"""One-off adapter: normalizes the real Samsung Experience Store master
dataset (Korean headers, vendor-specific shape) into the column contract
POST /api/import/gbp-urls expects (English snake_case).

This is exactly the kind of adapter a real integration needs: a source file
almost never matches your internal schema exactly. Run once:
  python convert_real_dataset.py
"""
import re
from pathlib import Path
from openpyxl import Workbook, load_workbook

here = Path(__file__).parent
src = here / "Samsung_Experience_Store_Global_Master_Dataset.xlsx"
wb_in = load_workbook(src, data_only=True)
ws_in = wb_in[wb_in.sheetnames[0]]

rows = list(ws_in.iter_rows(min_row=2, values_only=True))  # skip header

def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s in ("", "-") else s

def country_code(country_raw):
    m = re.search(r"\(([A-Za-z]{2,3})\)", country_raw or "")
    return m.group(1) if m else None

wb_out = Workbook()
ws_out = wb_out.active
ws_out.title = "gbp_urls"
headers = [
    "store_name", "region", "country", "gbp_url", "category",
    "region_group", "country_code", "address", "phone",
    "operating_status", "website_url", "rating", "review_count",
]
ws_out.append(headers)

for row in rows:
    (_no, region_group, country, city, store_name, address, phone,
     rating, review_count, status, maps_url, website_url) = row
    if not store_name or not maps_url:
        continue
    ws_out.append([
        store_name,
        city,
        country,
        maps_url,
        "Samsung Experience Store",
        region_group,
        country_code(country),
        address,
        clean(phone),
        status,
        clean(website_url),
        rating,
        int(review_count) if review_count is not None else None,
    ])

out = here / "gbp-urls.real.xlsx"
wb_out.save(out)
print(f"wrote {out} ({ws_out.max_row - 1} stores)")
